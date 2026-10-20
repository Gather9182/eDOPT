import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook
from pydantic import ValidationError

from app.models import ChargingCurve
from app.services.charging_curve import max_curve_energy
from app.services.feasibility import energy_feasibility_issues
from app.services.optimization import run_optimization
from app.services.solver import run_milp_solver, run_extended_milp_solver, simulate_baseline_load
from app.services.solver_diagnostics import OptimizationFailure
from app.utils.excel_exporter import ExcelExporter
from test_charging import inputs, bus, site, run_offline


def curve_inputs(initial=0, capacity=1000, **changes):
    location = {**site(), 'charger_power_kw':150,'charger_capacity_kw':150,'grid_limit_kw':150}
    vehicle = bus('curve', max_battery_capacity_kwh=capacity, initial_soc_kwh=initial,
                  min_final_soc_kwh=0, max_power_kw=150, availability=[1,0,0,0],
                  site_availability={'D':[1,0,0,0]})
    settings = dict(electricity_prices=[-100]*4, charging_curve={'enabled':True})
    settings.update(changes)
    return inputs([vehicle], [location], **settings)


class ChargingCurveTests(unittest.TestCase):
    def solve(self, data, extended=False):
        return run_offline(run_extended_milp_solver(data,[]) if extended else run_milp_solver(data,[]))

    def test_compact_model_avoids_end_band_binaries_for_tapering_curves(self):
        from app.services import solver
        class Captured(Exception):
            pass
        def capture(problem,*args):
            counts.append(sum(variable.isBinary() and variable.name.startswith('curve_') for variable in problem.variables()))
            raise Captured()
        counts = []
        for first_slot in (True,False):
            data = curve_inputs()
            available = [1,0,0,0] if first_slot else [0,1,0,0]
            data.buses[0].availability=available
            data.buses[0].site_availability={'D':available}
            with patch.object(solver,'solve_soc_policy',capture), self.assertRaises(Captured):
                self.solve(data)
        # Fixed initial SoC needs no band decisions; later states need only
        # the three internal thresholds, rather than eight segment decisions.
        self.assertEqual(counts,[0,3])

    def test_default_is_disabled_and_legacy_power_is_unchanged(self):
        data = curve_inputs(charging_curve={'enabled':False})
        result = self.solve(data)
        self.assertAlmostEqual(result.schedules['curve'][0],150)
        self.assertFalse(inputs([bus('old')]).charging_curve.enabled)

    def test_each_band_and_exact_boundaries_in_both_solvers(self):
        for extended in (False,True):
            for initial, expected in [(0,135),(600,97.5),(800,52.5),(900,22.5),(1000,0)]:
                with self.subTest(extended=extended,initial=initial):
                    data = curve_inputs(initial)
                    if extended:
                        data.scenario = 'pv'
                        data.energy_site_id = 'D'
                    result = self.solve(data,extended)
                    self.assertAlmostEqual(result.schedules['curve'][0],expected,places=3)

    def test_crossing_two_thresholds_integrates_time_instead_of_using_start_rate(self):
        result = self.solve(curve_inputs(initial=59,capacity=100))
        expected = 80 + (.25 - 1/135 - 20/97.5)*52.5
        self.assertAlmostEqual(result.soc_profiles['curve'][1],expected,places=3)
        self.assertLess(result.schedules['curve'][0],135)

    def test_non_monotone_custom_bands_and_bus_efficiency(self):
        data = curve_inputs(initial=49,capacity=100,charging_curve=dict(enabled=True,
            bands=[dict(end_soc_fraction=.5,power_fraction=.2),dict(end_soc_fraction=1,power_fraction=1)]))
        data.charging_sites[0].charger_power_kw=100
        data.buses[0].efficiency=.8
        expected = 50 + (.25 - 1/(20*.8))*100*.8
        result = self.solve(data)
        self.assertAlmostEqual(result.soc_profiles['curve'][1],expected,places=3)

    def test_lower_bus_and_connector_limits_are_respected_within_bands(self):
        for cap in (20,30,100):
            data = curve_inputs(initial=59,capacity=100)
            data.charging_sites[0].charger_power_kw=cap
            data.buses[0].max_power_kw=cap
            expected = 59 + max_curve_energy(data,data.buses[0],59,cap,.25,cap)
            result = self.solve(data)
            self.assertAlmostEqual(result.soc_profiles['curve'][1],expected,places=3)
            self.assertLessEqual(result.schedules['curve'][0],cap+.001)

    def test_different_chargers_use_their_own_percentage_base_in_both_solvers(self):
        for extended in (False,True):
            for initial, expected in [(0,[45,135]),(900,[7.5,22.5])]:
                with self.subTest(extended=extended,initial=initial):
                    vehicles = [bus(name,location,max_battery_capacity_kwh=1000,initial_soc_kwh=initial,
                                    min_final_soc_kwh=0,max_power_kw=200,availability=[1,0,0,0],
                                    site_availability={location:[1,0,0,0]}) for name,location in [('a','A'),('b','B')]]
                    locations = [{**site(location),'charger_power_kw':power,'charger_capacity_kw':500,'grid_limit_kw':500}
                                 for location,power in [('A',50),('B',150)]]
                    data = inputs(vehicles,locations,charging_curve={'enabled':True},electricity_prices=[-100]*4,
                                  scenario='pv' if extended else 'baseline',energy_site_id='A')
                    result = self.solve(data,extended)
                    for name,power in zip(('a','b'),expected):
                        self.assertAlmostEqual(result.schedules[name][0],power,places=3)
                    reference,_ = simulate_baseline_load(data,with_issues=True)
                    self.assertAlmostEqual(reference[0],sum(expected),places=3)
                    for vehicle,power in zip(data.buses,expected):
                        vehicle.min_final_soc_kwh=initial+power/4
                    self.assertEqual(energy_feasibility_issues(data),[])
                    data.buses[0].min_final_soc_kwh+=.5
                    self.assertTrue(energy_feasibility_issues(data))

    def test_bus_cap_is_applied_after_charger_percentage(self):
        data = curve_inputs()
        data.buses[0].max_power_kw=30
        result = self.solve(data)
        self.assertAlmostEqual(result.schedules['curve'][0],30)

    def test_legacy_mode_uses_bus_power_without_a_configured_location(self):
        data = curve_inputs(charging_sites=[])
        data.buses[0].max_power_kw=50
        self.assertAlmostEqual(self.solve(data).schedules['curve'][0],45)

    def test_reference_and_optimistic_precheck_use_the_same_curve(self):
        data = curve_inputs(initial=900)
        data.buses[0].min_final_soc_kwh=920
        self.assertTrue(energy_feasibility_issues(data))
        with self.assertRaisesRegex(RuntimeError,'Infeasible'):
            self.solve(data)
        load, issues = simulate_baseline_load(data,with_issues=True)
        self.assertAlmostEqual(load[0],22.5)
        self.assertTrue(issues)

    def test_lower_grid_limit_applies_during_band_crossing(self):
        data = curve_inputs(initial=89,capacity=100)
        data.charging_sites[0].grid_limit_kw=30
        result = self.solve(data)
        expected = 90 + (.25 - 1/30)*22.5
        self.assertAlmostEqual(result.soc_profiles['curve'][1],expected,places=3)

    def test_no_heuristic_fallback_when_curve_is_enabled_without_sites(self):
        data = curve_inputs(charging_sites=[])
        with patch('app.services.optimization.run_scenario_optimization',side_effect=RuntimeError('Test failure')):
            with patch('app.services.optimization.run_placeholder_optimization') as fallback:
                with self.assertRaises(OptimizationFailure):
                    run_offline(run_optimization(data))
                fallback.assert_not_called()

    def test_curve_settings_are_saved_and_exported(self):
        result = self.solve(curve_inputs())
        self.assertTrue(result.charging_curve['enabled'])
        self.assertEqual(result.charging_curve['power_basis'],'site_charger_power')
        self.assertNotIn('nominal_power_kw',result.charging_curve)
        book = load_workbook(BytesIO(ExcelExporter.generate_excel(result.model_dump()).getvalue()))
        self.assertIn('Charging Curve',book.sheetnames)

    def test_cyclic_soc_and_all_scenarios_retain_the_curve(self):
        from test_cyclic_soc import cycle_inputs, cycle_bus
        from app.services.solver import run_scenario_optimization
        for scenario in ('baseline','pv','bess','pv_bess'):
            with self.subTest(scenario=scenario):
                data = cycle_inputs([cycle_bus(initial=19,consumption=2)],scenario=scenario,
                                    energy_site_id='D',charging_curve={'enabled':True})
                result = run_offline(run_scenario_optimization(data,[]))
                self.assertAlmostEqual(result.soc_profiles['a'][-1],19,places=3)
                self.assertTrue(result.charging_curve['enabled'])
                self.assertTrue(result.end_condition['buses'][0]['compliant'])

    def test_additional_charging_visits_obey_curve_and_energy_balance(self):
        from test_opportunity_charging import data as visit_inputs, visiting_bus
        data = visit_inputs([visiting_bus()],charging_curve={'enabled':True})
        result = self.solve(data)
        self.assertEqual(len(result.opportunity_charging_visits),1)
        schedule = result.schedules['a']
        for t,power in enumerate(schedule):
            maximum = max_curve_energy(data,data.buses[0],result.soc_profiles['a'][t],40,.25)
            self.assertLessEqual(power*.25,maximum+.001)
        self.assertAlmostEqual(10+sum(schedule)/4-19,result.soc_profiles['a'][-1],places=3)

    def test_invalid_curves_are_rejected(self):
        for values in [dict(bands=[]),dict(bands=[dict(end_soc_fraction=.9,power_fraction=.5)]),
                       dict(bands=[dict(end_soc_fraction=1,power_fraction=0)]),
                       dict(bands=[dict(end_soc_fraction=.8,power_fraction=.5),dict(end_soc_fraction=.7,power_fraction=.4),dict(end_soc_fraction=1,power_fraction=.2)])]:
            with self.subTest(values=values), self.assertRaises(ValidationError):
                ChargingCurve(enabled=True,**values)
