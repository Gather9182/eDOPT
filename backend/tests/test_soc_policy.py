import unittest
from unittest.mock import patch
from io import BytesIO

from openpyxl import load_workbook
from pydantic import ValidationError

from app.models import OptimizationInputs
from app.main import validate_inputs
from app.services.solver import run_milp_solver, run_extended_milp_solver
from app.services.optimization import run_optimization
from app.services.solver_diagnostics import OptimizationFailure
from app.utils.excel_exporter import ExcelExporter
from app.utils.excel_processor import clean_distance
from test_charging import inputs, bus, run_offline, workbook
from test_opportunity_charging import data as visit_inputs, visiting_bus


class OperatingSoCTests(unittest.TestCase):
    def solve(self, data, extended=False):
        return run_offline(run_extended_milp_solver(data, []) if extended else run_milp_solver(data, []))

    def test_old_request_defaults_preserve_disabled_limits(self):
        result = self.solve(inputs([bus('old')]))
        self.assertEqual(result.soc_policy, {})
        self.assertAlmostEqual(result.soc_profiles['old'][-1], 20)

    def test_older_enabled_maximum_keeps_soft_semantics(self):
        data = inputs([bus('legacy-max', min_final_soc_kwh=18)], operating_max_soc_enabled=True,
                      operating_max_soc_fraction=.5)
        self.assertEqual(data.operating_max_soc_mode,'soft')
        self.assertAlmostEqual(self.solve(data).soc_profiles['legacy-max'][-1],18,places=4)

    def test_hard_max_limits_both_solver_paths_even_with_negative_prices(self):
        for extended in (False, True):
            with self.subTest(extended=extended):
                data = inputs([bus('hard-max',min_final_soc_kwh=2)],electricity_prices=[-100]*4,
                    operating_max_soc_enabled=True,operating_max_soc_mode='hard',operating_max_soc_fraction=.5)
                result = self.solve(data,extended)
                self.assertAlmostEqual(max(result.soc_profiles['hard-max']),10,places=4)
                self.assertEqual(result.soc_policy['max_mode'],'hard')
                self.assertIsNone(result.soc_policy['minimum_deviation_kwh_h'])
                self.assertTrue(result.baseline_feasible)

    def test_hard_max_rejects_conflicting_end_target_while_soft_max_can_relax(self):
        data = inputs([bus('end-target',min_final_soc_kwh=18)],operating_max_soc_enabled=True,
                      operating_max_soc_mode='hard',operating_max_soc_fraction=.5)
        checked = run_offline(validate_inputs(data))
        self.assertIn('BUS_FINAL_TARGET_EXCEEDS_OPERATING_MAX',[issue.get('code') for issue in checked['issues']])
        for extended in (False,True):
            with self.subTest(extended=extended),self.assertRaises(RuntimeError):
                self.solve(data,extended)
        data.operating_max_soc_mode='soft'
        self.assertAlmostEqual(self.solve(data).soc_profiles['end-target'][-1],18,places=4)

    def test_hard_max_checks_initial_state_and_disabled_maximum_has_no_effect(self):
        data = inputs([bus('initial-max',initial_soc_kwh=20,min_final_soc_kwh=0)],
            operating_max_soc_enabled=True,operating_max_soc_mode='hard',operating_max_soc_fraction=.5)
        checked = run_offline(validate_inputs(data))
        self.assertIn('BUS_OPERATING_MAX_INITIAL',[issue.get('code') for issue in checked['issues']])
        with self.assertRaises(RuntimeError):
            self.solve(data)
        data.operating_max_soc_enabled=False
        self.assertEqual(self.solve(data).soc_profiles['initial-max'],[20]*5)

    def test_hard_max_applies_to_optional_charging_trips(self):
        for extended in (False,True):
            with self.subTest(extended=extended):
                data = visit_inputs([visiting_bus()],operating_max_soc_enabled=True,
                    operating_max_soc_mode='hard',operating_max_soc_fraction=.9)
                result = self.solve(data,extended)
                self.assertTrue(result.opportunity_charging_visits)
                self.assertLessEqual(max(result.soc_profiles['a']),18.0001)
                data.operating_max_soc_fraction=.7
                with self.assertRaises(RuntimeError):
                    self.solve(data,extended)

    def test_invalid_max_mode_is_rejected(self):
        with self.assertRaises(ValidationError):
            OptimizationInputs(**{**inputs([bus('invalid-mode')]).model_dump(),'operating_max_soc_mode':'unknown'})

    def test_soft_max_prevents_cheap_surplus_even_at_negative_prices(self):
        for extended in (False, True):
            with self.subTest(extended=extended):
                data = inputs([bus('max', min_final_soc_kwh=2)], electricity_prices=[-100]*4,
                    operating_max_soc_enabled=True, operating_max_soc_fraction=.5)
                result = self.solve(data, extended)
                self.assertAlmostEqual(max(result.soc_profiles['max']), 10, places=4)
                self.assertEqual(result.soc_policy['violations'], [])

    def test_soft_max_relaxes_only_as_much_as_required_by_end_target(self):
        data = inputs([bus('necessary', min_final_soc_kwh=18)], operating_max_soc_enabled=True,
                      operating_max_soc_fraction=.5, electricity_prices=[-100]*4)
        result = self.solve(data)
        self.assertAlmostEqual(result.soc_profiles['necessary'][-1], 18, places=4)
        self.assertTrue(result.soc_policy['violations'])
        self.assertAlmostEqual(result.soc_policy['violations'][-1]['deviation_kwh'], 8)

    def test_soft_min_accepts_unavoidable_dip_and_reports_actual_time(self):
        data = inputs([bus('dip', initial_soc_kwh=20, min_final_soc_kwh=0,
            availability=[0,1,1,1], site_availability={'D':[0,1,1,1]}, trip_energy_profile_kwh=[19,0,0,0])],
            operating_min_soc_mode='soft', operating_min_soc_fraction=.1, horizon_start_minute=1440)
        result = self.solve(data)
        violation = result.soc_policy['violations'][0]
        self.assertEqual(violation['bus_id'], 'dip')
        self.assertEqual(violation['kind'], 'below_min')
        self.assertEqual(violation['minute'], 1455)
        self.assertAlmostEqual(violation['deviation_kwh'], 1)
        self.assertTrue(result.baseline_feasible)

    def test_hard_min_rejects_same_route_and_validation_catches_it(self):
        data = inputs([bus('dip', initial_soc_kwh=20, min_final_soc_kwh=0,
            availability=[0,1,1,1], site_availability={'D':[0,1,1,1]}, trip_energy_profile_kwh=[19,0,0,0])],
            operating_min_soc_mode='hard', operating_min_soc_fraction=.1)
        self.assertEqual(run_offline(validate_inputs(data))['status'], 'invalid')
        with self.assertRaises(RuntimeError):
            self.solve(data)

    def test_soft_min_is_not_sacrificed_for_cheaper_later_charging(self):
        data = inputs([bus('reserve', initial_soc_kwh=2, min_final_soc_kwh=0,
            trip_energy_profile_kwh=[2,0,0,0])], operating_min_soc_mode='soft',
            electricity_prices=[1000,1,1,1])
        result = self.solve(data)
        self.assertGreaterEqual(result.soc_profiles['reserve'][1], 1.999)
        self.assertGreater(result.schedules['reserve'][0], 7.99)
        self.assertEqual(result.soc_policy['violations'], [])

    def test_initial_high_soc_is_not_artificially_discharged(self):
        data = inputs([bus('initial', initial_soc_kwh=20, min_final_soc_kwh=0)],
            operating_max_soc_enabled=True, operating_max_soc_fraction=.9)
        result = self.solve(data)
        self.assertEqual(result.soc_profiles['initial'], [20]*5)
        self.assertTrue(result.soc_policy['violations'][0]['initial_state'])
        self.assertEqual(sum(result.schedules['initial']), 0)

    def test_initial_below_hard_min_is_invalid_but_soft_min_is_allowed(self):
        data = inputs([bus('initial', initial_soc_kwh=1)], operating_min_soc_mode='hard')
        self.assertEqual(run_offline(validate_inputs(data))['status'], 'invalid')
        data.operating_min_soc_mode = 'soft'
        self.assertEqual(self.solve(data).soc_policy['violations'][0]['state_index'], 0)

    def test_soft_modes_never_relax_physical_zero_or_capacity(self):
        data = inputs([bus('physical', initial_soc_kwh=1, min_final_soc_kwh=0,
            availability=[0,1,1,1], site_availability={'D':[0,1,1,1]}, trip_energy_profile_kwh=[3,0,0,0])],
            operating_min_soc_mode='soft', operating_max_soc_enabled=True)
        with self.assertRaises(RuntimeError):
            self.solve(data)

    def test_both_limits_include_selected_optional_trip_energy(self):
        data = visit_inputs([visiting_bus()], operating_min_soc_mode='soft',
            operating_max_soc_enabled=True, operating_max_soc_fraction=.7)
        result = self.solve(data)
        self.assertTrue(result.opportunity_charging_visits)
        self.assertEqual(sum(result.additional_trip_energy_profiles['a']),4)
        self.assertTrue(all(0 <= value <= 20 for value in result.soc_profiles['a']))
        self.assertAlmostEqual(10 + sum(result.schedules['a']) / 4 - 19, result.soc_profiles['a'][-1], places=3)

    def test_unproven_primary_objective_cannot_fall_through_to_cost_optimization(self):
        data = inputs([bus('proof')],operating_min_soc_mode='soft')
        with patch('app.services.soc_policy.solve_with_diagnostics') as mocked:
            with self.assertRaisesRegex(OptimizationFailure,'not proven'):
                self.solve(data)
            self.assertEqual(mocked.call_count,1)

    def test_no_heuristic_fallback_with_operating_limits(self):
        data = inputs([bus('fallback')], operating_min_soc_mode='soft')
        data.charging_sites = []
        with patch('app.services.optimization.run_scenario_optimization', side_effect=RuntimeError('test failure')):
            with self.assertRaises(OptimizationFailure):
                run_offline(run_optimization(data))

    def test_invalid_band_and_nonfinite_values_are_rejected(self):
        original = inputs([bus('valid')]).model_dump()
        for patch_values in [dict(operating_min_soc_fraction=-.1), dict(operating_max_soc_fraction=1.1),
                             dict(operating_min_soc_fraction=float('nan')),
                             dict(operating_min_soc_mode='soft', operating_min_soc_fraction=.95, operating_max_soc_enabled=True)]:
            with self.subTest(values=patch_values), self.assertRaises(ValidationError):
                OptimizationInputs(**{**original, **patch_values})

    def test_reference_is_invalid_when_final_energy_target_is_not_reached_without_sites(self):
        data = inputs([bus('reference', min_final_soc_kwh=20, initial_soc_kwh=20,
                            trip_energy_profile_kwh=[0,0,0,10])])
        data.charging_sites = []
        result = self.solve(data)
        self.assertFalse(result.baseline_feasible)

    def test_excel_includes_policy_and_no_savings_for_invalid_reference(self):
        data = inputs([bus('export', min_final_soc_kwh=18)], operating_max_soc_enabled=True,
                      operating_max_soc_fraction=.5)
        result = self.solve(data).model_dump()
        result['baseline_feasible'] = False
        result['run_metadata'] = {'inputs':data.model_dump()}
        workbook = load_workbook(BytesIO(ExcelExporter.generate_excel(result).getvalue()))
        self.assertIn('Operating SoC Deviations', workbook.sheetnames)
        values = [cell.value for sheet in workbook for row in sheet for cell in row]
        self.assertTrue(any(isinstance(value,str) and 'n.a. (reference not comparable)' in value for value in values))
        self.assertTrue(any(isinstance(value,str) and 'planning hours' in value for value in values))


class StrictDistanceTests(unittest.TestCase):
    def test_complete_import_retains_source_rows_and_recomputes_all_energy(self):
        from app.utils.excel_processor import process_umlauf_excel
        source = workbook([dict(Umlauf=f'R{i % 8}',Meter=value) for i,value in enumerate([21]*20+[25,35,94])])
        legacy = process_umlauf_excel(source,dict(BETRIEBSTAGE=[],DISTANCE_UNIT='legacy'))
        mixed = process_umlauf_excel(source,dict(BETRIEBSTAGE=[],DISTANCE_UNIT='mixed',MIXED_KM_THRESHOLD=50))
        legacy_km = sum(trip['dist_m'] for bus in legacy['bus_inputs'] for trip in bus['trips']) / 1000
        mixed_km = sum(trip['dist_m'] for bus in mixed['bus_inputs'] for trip in bus['trips']) / 1000
        self.assertAlmostEqual(mixed_km - legacy_km,479.52)
        warning = next(row for row in legacy['validation_results']['warnings'] if row['type']=='DISTANCE_UNIT')
        self.assertEqual(warning['count'],23)
        self.assertTrue(warning['examples'])
        trip = mixed['bus_inputs'][0]['trips'][0]
        self.assertEqual(trip['source_excel_row'],2)
        self.assertEqual(trip['distance_original'],'21')
        self.assertEqual(trip['distance_unit_rule'],'mixed')
        self.assertEqual(trip['dist_m'],21000)

    def test_invalid_source_identifies_sheet_row_and_route(self):
        from app.utils.excel_processor import process_umlauf_excel
        with self.assertRaisesRegex(ValueError,'Excel row 2, route U'):
            process_umlauf_excel(workbook([dict(Meter='broken')]),dict(BETRIEBSTAGE=[]))

    def test_explicit_units_and_known_mixed_examples(self):
        for source, expected in [(21,21000),(25,25000),(35,35000),(94,94),('1,5 km',1500),('94 m',94)]:
            self.assertEqual(clean_distance(source,'mixed',50),expected)
        self.assertEqual(clean_distance('1.5 km','legacy'),1500)
        self.assertEqual(clean_distance('21,5','legacy'),21500)

    def test_invalid_values_do_not_become_zero(self):
        for mode in ('legacy','m','km','mixed'):
            for value in (None,'','not a distance',-1,float('inf'),float('nan')):
                with self.subTest(mode=mode,value=value), self.assertRaises(ValueError):
                    clean_distance(value,mode)
        self.assertEqual(clean_distance(0),0)
