import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook
from pydantic import ValidationError

from app.models import OptimizationInputs
from app.main import validate_inputs
from app.services.optimization import run_optimization
from app.services.solver import run_milp_solver, run_extended_milp_solver, run_scenario_optimization
from app.services.solver_diagnostics import OptimizationFailure
from app.services.route_feasibility import classify_routes
from app.utils.excel_exporter import ExcelExporter
from test_charging import inputs, bus, run_offline
from test_opportunity_charging import visiting_bus


def cycle_bus(name='a', initial=10, consumption=6, **changes):
    values = dict(initial_soc_kwh=initial,min_final_soc_kwh=2,
                  trip_energy_profile_kwh=[0]*4+[consumption]+[0]*91)
    values.update(changes)
    return bus(name,n=96,**values)


def cycle_inputs(buses=None, **changes):
    values = dict(planning_horizon_hours=24,electricity_prices=[100]*96,
                  cyclic_connections=True,use_cyclic_soc=True)
    values.update(changes)
    return inputs(buses or [cycle_bus()],**values)


class CyclicBusSoCTests(unittest.TestCase):
    def test_all_scenarios_enforce_equality_per_bus_not_fleet_balance(self):
        for scenario in ('baseline','pv','bess','pv_bess'):
            with self.subTest(scenario=scenario):
                source=cycle_inputs([cycle_bus('a',10,6),cycle_bus('b',8,4)],scenario=scenario,energy_site_id='D')
                result=run_offline(run_scenario_optimization(source,[]))
                for identifier,initial in [('a',10),('b',8)]:
                    self.assertAlmostEqual(result.soc_profiles[identifier][-1],initial,places=3)
                self.assertTrue(all(row['compliant'] for row in result.end_condition['buses']))
                self.assertEqual(result.end_condition['mode'],'cyclic_soc')
                self.assertFalse(result.end_condition['next_day_readiness_enabled'])

    def test_negative_prices_cannot_add_surplus_to_cyclic_bus(self):
        source=cycle_inputs(electricity_prices=[-100]*96)
        result=run_offline(run_milp_solver(source))
        self.assertAlmostEqual(sum(result.schedules['a'])/4,6,places=3)
        self.assertAlmostEqual(result.soc_profiles['a'][-1],10,places=3)

    def test_chosen_optional_round_trip_energy_is_replaced_before_end(self):
        original=visiting_bus(initial_soc_kwh=2)
        for key in ('availability','trip_energy_profile_kwh'):
            original[key]=original[key]+[0]*88
        for key in ('site_availability','site_presence','site_arrivals'):
            original[key]={site:values+[0]*88 for site,values in original[key].items()}
        result=run_offline(run_milp_solver(cycle_inputs([original])))
        self.assertTrue(result.opportunity_charging_visits)
        self.assertEqual(sum(result.additional_trip_energy_profiles['a']),4)
        self.assertAlmostEqual(sum(result.schedules['a'])/4,19,places=3)
        self.assertAlmostEqual(result.soc_profiles['a'][-1],2,places=3)

    def test_readiness_requirement_is_replaced_not_added(self):
        result=run_offline(run_milp_solver(cycle_inputs([cycle_bus(next_day_required_energy_kwh=100)])))
        self.assertAlmostEqual(result.soc_profiles['a'][-1],10,places=3)
        self.assertEqual(result.bus_diagnostics[0].target_min_final_soc_kwh,10)
        self.assertEqual(result.bus_diagnostics[0].next_day_required_energy_kwh,0)

    def test_reserve_above_initial_is_reported_and_never_silently_dropped(self):
        source=cycle_inputs([cycle_bus(min_final_soc_kwh=12)])
        checked=run_offline(validate_inputs(source))
        self.assertEqual(checked['status'],'invalid')
        self.assertIn('BUS_CYCLIC_RESERVE_CONFLICT',[row.get('code') for row in checked['issues']])
        with self.assertRaises(RuntimeError):
            run_offline(run_milp_solver(source))

    def test_energy_check_uses_initial_energy_as_cyclic_end_target(self):
        source=cycle_inputs([cycle_bus(max_power_kw=.1)])
        checked=run_offline(validate_inputs(source))
        self.assertTrue(any(row.get('code')=='BUS_ENERGY_INFEASIBLE' for row in checked['issues']))

    def test_cyclic_failure_never_uses_legacy_heuristic_even_without_sites(self):
        source=cycle_inputs()
        source.charging_sites=[]
        with patch('app.services.optimization.run_scenario_optimization',side_effect=RuntimeError('test failure')):
            with self.assertRaises(OptimizationFailure):
                run_offline(run_optimization(source))

    def test_reference_must_meet_equality_not_only_lower_end_reserve(self):
        result=run_offline(run_milp_solver(cycle_inputs()))
        self.assertFalse(result.baseline_feasible)
        self.assertTrue(any('Start SoC = End SoC' in issue for issue in result.baseline_issues))

    def test_bess_does_not_receive_cyclic_bus_constraint(self):
        source=cycle_inputs(scenario='bess',energy_site_id='D',bess_config=dict(capacity_kwh=20,
            initial_soc_kwh=5,min_soc_kwh=0,max_soc_kwh=20,max_charge_power_kw=10,max_discharge_power_kw=10,efficiency=1))
        result=run_offline(run_extended_milp_solver(source,[]))
        self.assertAlmostEqual(result.soc_profiles['a'][-1],10,places=3)
        self.assertAlmostEqual(result.bess_soc_profile[-1],0,places=3)

    def test_infeasible_grid_only_comparison_does_not_block_feasible_pv_cycle(self):
        source=cycle_inputs(scenario='pv',grid_limit_kw=.1,
            pv_config=dict(capacity_kw=40,profile_kwh=[10]*96))
        source.charging_sites=[]
        result=run_offline(run_scenario_optimization(source,[]))
        self.assertIsNone(result.basic_optimized_cost_eur)
        self.assertAlmostEqual(result.soc_profiles['a'][-1],10,places=3)

    def test_cyclic_soc_is_independent_of_disabled_hard_or_soft_operating_limits(self):
        for mode in ('hard','soft'):
            with self.subTest(mode=mode):
                result=run_offline(run_milp_solver(cycle_inputs(operating_min_soc_mode=mode,
                    operating_max_soc_enabled=True,operating_max_soc_mode=mode,operating_max_soc_fraction=.9)))
                self.assertAlmostEqual(result.soc_profiles['a'][-1],10,places=3)

    def test_exact_isolated_diagnostic_retains_cyclic_constraint(self):
        source=cycle_inputs([cycle_bus('unchargeable',max_power_kw=0)])
        diagnostic=run_offline(classify_routes(source,budget_seconds=10))
        self.assertEqual(diagnostic['routes'][0]['status'],'individually_infeasible')

    def test_export_contains_actual_cycle_checks_and_settings(self):
        source=cycle_inputs()
        result=run_offline(run_milp_solver(source)).model_dump()
        result['run_metadata']={'inputs':source.model_dump()}
        workbook=load_workbook(BytesIO(ExcelExporter.generate_excel(result).getvalue()))
        self.assertIn('End Condition Settings',workbook.sheetnames)
        self.assertIn('Bus End Conditions',workbook.sheetnames)
        self.assertEqual(workbook['Bus End Conditions']['D2'].value,0)

    def test_old_requests_keep_both_old_readiness_modes(self):
        for readiness,target in [(False,2),(True,6)]:
            with self.subTest(readiness=readiness):
                source=cycle_inputs([cycle_bus(initial=10,consumption=6,next_day_required_energy_kwh=4)],
                    use_cyclic_soc=False,use_next_day_readiness=readiness)
                result=run_offline(run_milp_solver(source))
                self.assertAlmostEqual(result.soc_profiles['a'][-1],max(4,target),places=3)
                self.assertFalse(result.end_condition['cyclic_soc_enabled'])

    def test_direct_api_rejects_conflicting_flags_horizons_and_chunks(self):
        original=cycle_inputs().model_dump()
        for values in [dict(use_next_day_readiness=True),dict(planning_horizon_hours=48),
                       dict(electricity_prices=[100]*4),dict(cyclic_connections=False)]:
            with self.subTest(values=values),self.assertRaises(ValidationError):
                OptimizationInputs(**{**original,**values})
